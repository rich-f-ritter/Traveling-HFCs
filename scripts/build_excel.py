#!/usr/bin/env python3
"""Build the sortable, styled master Excel workbook from data/census_master.json.

Output: data/Traveling_HFC_PFC_Master.xlsx
Sheets: Master List (autofilter + frozen header + tier color bands), Summary, Legend.
"""
import json
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
DATA = json.loads((ROOT / "data/census_master.json").read_text())

# (header, source-key, width, kind)  kind: text|num|money|year|wrap
COLUMNS = [
    ("Census ID", "census_id", 14, "text"),
    ("Property Name", "name", 30, "text"),
    ("Address", "address", 26, "text"),
    ("City", "city", 15, "text"),
    ("County", "county", 12, "text"),
    ("Units", "units", 8, "num"),
    ("Year Built", "year_built", 10, "year"),
    ("Sponsor (verified)", "sponsor", 34, "text"),
    ("Type", "species", 7, "text"),
    ("Statute", "statute", 9, "text"),
    ("Tier", "tier", 6, "num"),
    ("Traveling", "traveling", 10, "text"),
    ("Faces 2027 Cliff", "_cliff", 15, "text"),
    ("Kill Date", "kill_date", 12, "text"),
    ("Kill Regime", "kill_regime", 22, "wrap"),
    ("Exemption Status", "exemption_status", 15, "text"),
    ("Appraised Value", "appraised_value", 16, "money"),
    ("First Exempt Yr", "first_exempt_year", 12, "year"),
    ("Private Partner — Group", "private_partner_parent", 26, "text"),
    ("Private Partner — Seller SPE", "private_partner_spe", 30, "text"),
    ("Lender", "lender", 26, "text"),
    ("Loan Amount", "lender_loan_amount", 15, "money"),
    ("Recent Capital Event", "recent_capital_event", 24, "wrap"),
    ("Litigation", "litigation", 30, "wrap"),
    ("CAD Account", "cad_account", 18, "text"),
    ("Confidence", "confidence", 13, "text"),
    ("Partner Source", "private_partner_source", 28, "wrap"),
    ("Lender Source", "lender_source", 28, "wrap"),
    ("Year Built Source", "year_built_source", 22, "wrap"),
    ("Kill Basis / Notes", "kill_basis", 40, "wrap"),
]

TIER_FILL = {
    1: PatternFill("solid", fgColor="FCE4E4"),   # light red — going away 2027
    2: PatternFill("solid", fgColor="FFF6DA"),   # light amber — grandfathered PFC
    3: PatternFill("solid", fgColor="EDEDED"),   # light gray — housing authority
}
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def money(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v)
    m = re.search(r"([\d,]+(?:\.\d+)?)\s*(m|mm|million|k|b|billion)?", s, re.I)
    if not m:
        return None
    num = float(m.group(1).replace(",", ""))
    unit = (m.group(2) or "").lower()
    if unit in ("m", "mm", "million"):
        num *= 1e6
    elif unit in ("b", "billion"):
        num *= 1e9
    elif unit == "k":
        num *= 1e3
    return num


def as_int(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def sort_key(r):
    t = r.get("tier") or 9
    v = r.get("appraised_value")
    v = v if isinstance(v, (int, float)) else 0
    return (t, -v)


def build_master(wb):
    ws = wb.active
    ws.title = "Master List"
    # header
    for c, (hdr, *_ ) in enumerate(COLUMNS, 1):
        cell = ws.cell(1, c, hdr)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = BORDER
    rows = sorted(DATA, key=sort_key)
    for ri, r in enumerate(rows, 2):
        tier = r.get("tier")
        for ci, (hdr, key, width, kind) in enumerate(COLUMNS, 1):
            if key == "_cliff":
                val = "YES — 1/1/2027" if tier == 1 else "No"
            else:
                val = r.get(key)
            if kind == "money":
                val = money(val)
            elif kind in ("num", "year"):
                val = as_int(val)
            cell = ws.cell(ri, ci, val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", wrap_text=(kind == "wrap"))
            if kind == "money" and isinstance(val, float):
                cell.number_format = '#,##0'
            if TIER_FILL.get(tier) and ci <= 13:
                cell.fill = TIER_FILL[tier]
    # widths, filter, freeze
    for ci, (hdr, key, width, kind) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(ci)].width = width
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{len(rows)+1}"
    ws.freeze_panes = "C2"   # freeze header row + first two cols (id, name)
    ws.row_dimensions[1].height = 30
    ws.sheet_view.zoomScale = 90


def money_str(x):
    return f"${x/1e9:.2f}B" if x >= 1e9 else f"${x/1e6:.1f}M"


def build_summary(wb):
    from collections import defaultdict
    ws = wb.create_sheet("Summary")
    by_tier = defaultdict(lambda: [0, 0.0])
    by_sponsor = defaultdict(lambda: [0, 0.0])
    by_county = defaultdict(lambda: [0, 0.0])
    total = 0.0
    for r in DATA:
        v = r.get("appraised_value")
        v = v if isinstance(v, (int, float)) else 0
        total += v
        by_tier[r.get("tier")][0] += 1; by_tier[r.get("tier")][1] += v
        by_sponsor[r.get("sponsor")][0] += 1; by_sponsor[r.get("sponsor")][1] += v
        by_county[r.get("county")][0] += 1; by_county[r.get("county")][1] += v

    def block(title, d, r0, keyorder=None):
        ws.cell(r0, 1, title).font = Font(bold=True, size=12)
        ws.cell(r0 + 1, 1, "Group").font = Font(bold=True)
        ws.cell(r0 + 1, 2, "Properties").font = Font(bold=True)
        ws.cell(r0 + 1, 3, "Appraised Value").font = Font(bold=True)
        items = sorted(d.items(), key=lambda kv: -kv[1][1]) if keyorder is None else \
            [(k, d[k]) for k in keyorder if k in d]
        rr = r0 + 2
        for k, (n, v) in items:
            ws.cell(rr, 1, str(k))
            ws.cell(rr, 2, n)
            c = ws.cell(rr, 3, round(v)); c.number_format = '#,##0'
            rr += 1
        return rr + 1

    ws.cell(1, 1, "Texas Traveling HFC / PFC Census — Summary").font = Font(bold=True, size=14)
    ws.cell(2, 1, f"{len(DATA)} properties · {money_str(total)} total appraised value · as of 2026-07-12")
    tier_lbl = {1: "Tier 1 — HFC (Ch.394), exemption dies 1/1/2027",
                2: "Tier 2 — PFC (Ch.303), grandfathered (no cliff)",
                3: "Tier 3 — Housing authority (Ch.392), no statutory cliff"}
    r = block("By Tier", {tier_lbl.get(k, k): v for k, v in by_tier.items()}, 4)
    r = block("By Sponsor", by_sponsor, r)
    block("By County", by_county, r)
    for ci, w in enumerate([46, 14, 20], 1):
        ws.column_dimensions[get_column_letter(ci)].width = w


def build_legend(wb):
    ws = wb.create_sheet("Legend & Notes")
    notes = [
        ("Texas Traveling HFC / PFC Census — Legend & Notes", True),
        ("", False),
        ("Scope: every Texas multifamily property held in a traveling public-corporation tax-exemption", False),
        ("structure. 458 properties across 12 counties (126 found beyond the original Yardi Matrix seed).", False),
        ("", False),
        ("TIERS (color bands on the Master List)", True),
        ("Tier 1 (pink) — HFC, Local Gov't Code Ch. 394. Exemption terminates 1/1/2027 under HB 21 §13(i)", False),
        ("   if HFC-owned out-of-jurisdiction on 9/1/2025 with no host resolutions; §13(e) kills it earlier", False),
        ("   on a sale/refi/majority transfer. THIS IS THE 'GOING AWAY' SET.", False),
        ("Tier 2 (amber) — PFC, Ch. 303. Grandfathered by HB 2071 — NO 2027 cliff; unwinds on capital", False),
        ("   events, audit failure, or CAD/AG challenge.", False),
        ("Tier 3 (gray) — Housing authority, Ch. 392. No statutory cliff; exposure via litigation only.", False),
        ("", False),
        ("KEY COLUMNS", True),
        ("Traveling — property sits outside the sponsor's home jurisdiction (city for city sponsors,", False),
        ("   county for county sponsors; first-order test, city+5mi edge cases need geocode).", False),
        ("Exemption Status — several North Texas CADs already stripped these exemptions ahead of 2027", False),
        ("   (e.g., Tarrant cut them ~Feb 2025); 'taxable' here means the property is paying full tax now.", False),
        ("Appraised Value — CAD appraised value; TX is a non-disclosure state, so this ≠ trade price.", False),
        ("Private Partner — the group that put the building into the HFC/PFC: Seller SPE = the deed grantor;", False),
        ("   Group = the parent developer/equity firm. Corrected against the Matrix seed where they conflicted.", False),
        ("Lender — from freely-viewable deeds of trust + SEC/CMBS filings + press only.", False),
        ("", False),
        ("COVERAGE (public-record only)", True),
        ("Year built 92% · Private partner (group) 78% · Seller SPE 68% · Lender 21%.", False),
        ("Lender is low by design: deed-of-trust records sit behind locked county-clerk portals, and", False),
        ("post-flip debt often rides the leasehold (invisible to fee-record lookups). Blank ≠ no loan.", False),
        ("", False),
        ("CAVEATS", True),
        ("~10 rows have no name/address/CAD account (upstream Matrix gaps). A few sponsor tags are flagged", False),
        ("for re-verification (see docs/enrichment_notes.md). Confidence column marks per-row certainty.", False),
        ("Full methodology: PROJECT_PLAN.md, docs/phase1_verdict.md, docs/enrichment_notes.md.", False),
    ]
    for i, (text, bold) in enumerate(notes, 1):
        c = ws.cell(i, 1, text)
        if bold:
            c.font = Font(bold=True, size=12 if i == 1 else 11)
    ws.column_dimensions["A"].width = 100


wb = Workbook()
build_master(wb)
build_summary(wb)
build_legend(wb)
out = ROOT / "data/Traveling_HFC_PFC_Master.xlsx"
wb.save(out)
print(f"wrote {out} ({len(DATA)} rows)")
